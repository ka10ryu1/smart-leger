"""household.xlsx を DB として読み書きするリポジトリ

保存手順:
    1. 同じフォルダの一時ファイルへ書き出す
    2. 一時ファイルを openpyxl で再オープンして壊れていないことを確認
    3. 既存の正本を世代バックアップ
    4. os.replace で正本を置換（同一ボリューム内なので実質 atomic）
    5. Dropbox が設定されていれば latest / backup へコピー

Excel が他アプリで開かれてロックされている場合は ExcelLockedError を投げる
"""

from __future__ import annotations

import logging
import os
from pathlib import Path
from typing import Any, Callable, Iterator

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..constants import DEFAULT_CATEGORIES, EXCEL_COLUMN_WIDTHS
from ..models import (
    Allocation,
    Category,
    ImportRecord,
    LedgerData,
    MerchantRule,
    Transaction,
)
from .backup import DropboxBackup, create_generation_backup

logger = logging.getLogger(__name__)


class ExcelLockedError(Exception):
    """Excel が他のアプリで開かれていて書き込めない"""


class ExcelSaveError(Exception):
    """一時ファイルの保存・検証に失敗した"""


def sheet_columns() -> dict[str, tuple[str, ...]]:
    """シート名 → 列名タプルの対応を返す"""
    return {
        'transactions': Transaction.COLUMNS,
        'merchant_rules': MerchantRule.COLUMNS,
        'categories': Category.COLUMNS,
        'imports': ImportRecord.COLUMNS,
        'allocations': Allocation.COLUMNS,
    }


def default_ledger() -> LedgerData:
    """初期カテゴリだけが入った空の LedgerData を返す"""
    data = LedgerData()
    data.categories = [Category(category=c, sort_order=i + 1) for i, c in enumerate(DEFAULT_CATEGORIES)]
    return data


def rows_as_dicts(ws: Worksheet, columns: tuple[str, ...]) -> Iterator[dict[str, Any]]:
    """シートの 2 行目以降を列名 → 値の dict として返す（欠けている列は None で補う）

    Args:
        ws: 読み取り対象のシート
        columns: 期待する列名（旧形式のファイルで足りない列を補うために使う）
    """
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return

    names = [str(h).strip() if h is not None else '' for h in header]
    for values in rows:
        record = {name: value for name, value in zip(names, values) if name}
        if record.get(columns[0]) in (None, ''):  # 先頭列（id / pattern など）が空の行は不完全なので読まない
            continue

        for col in columns:
            record.setdefault(col, None)

        yield record


class ExcelRepository:
    """household.xlsx の読み書きを担当する"""

    def __init__(
        self,
        excel_path: Path,
        backup_dir: Path | None = None,
        backup_generations: int = 20,
        dropbox: DropboxBackup | None = None,
    ):
        """
        Args:
            excel_path: 正本 Excel のパス
            backup_dir: 世代バックアップ先（None なら excel_path と同じ場所の backup/）
            backup_generations: 残す世代数
            dropbox: Dropbox コピー（None なら無効）
        """
        self.excel_path = Path(excel_path)
        self.backup_dir = Path(backup_dir) if backup_dir else self.excel_path.parent / 'backup'
        self.backup_generations = backup_generations
        self.dropbox = dropbox or DropboxBackup(None)
        self.last_dropbox_result: dict[str, Path] | None = None

    # ------------------------------------------------------------------ read
    def load(self) -> LedgerData:
        """全シートを読み込む（正本が無ければ初期カテゴリ入りの新規ファイルを作成する）"""
        if not self.excel_path.exists():
            logger.info('excel not found: creating=%s', self.excel_path)
            data = default_ledger()
            self.save(data)
            return data

        try:
            wb = load_workbook(self.excel_path, read_only=True, data_only=True)
        except PermissionError as exc:
            raise ExcelLockedError(self.locked_message()) from exc

        try:
            data = self.read_sheets(wb)
        finally:
            wb.close()

        if not data.categories:
            data.categories = default_ledger().categories

        return data

    def read_sheets(self, wb: Workbook) -> LedgerData:
        """ワークブックの各シートを LedgerData に変換する（存在しないシートは空扱い）

        Args:
            wb: 読み取り専用で開いたワークブック
        """
        data = LedgerData()
        columns = sheet_columns()
        if 'transactions' in wb.sheetnames:
            data.transactions = [
                Transaction.from_row(r) for r in rows_as_dicts(wb['transactions'], columns['transactions'])
            ]

        if 'merchant_rules' in wb.sheetnames:
            data.merchant_rules = [
                MerchantRule.from_row(r) for r in rows_as_dicts(wb['merchant_rules'], columns['merchant_rules'])
            ]

        if 'categories' in wb.sheetnames:
            data.categories = [Category.from_row(r) for r in rows_as_dicts(wb['categories'], columns['categories'])]

        if 'imports' in wb.sheetnames:
            data.imports = [ImportRecord.from_row(r) for r in rows_as_dicts(wb['imports'], columns['imports'])]

        if 'allocations' in wb.sheetnames:
            data.allocations = [
                Allocation.from_row(r) for r in rows_as_dicts(wb['allocations'], columns['allocations'])
            ]

        return data

    # ----------------------------------------------------------------- write
    def build_workbook(self, data: LedgerData) -> Workbook:
        """LedgerData から 5 シート構成のワークブックを組み立てる

        Args:
            data: 書き出すデータ
        """
        wb = Workbook()
        wb.remove(wb.worksheets[0])
        sheet_rows: dict[str, list[list[Any]]] = {
            'transactions': [t.to_row() for t in data.transactions],
            'merchant_rules': [r.to_row() for r in data.merchant_rules],
            'categories': [c.to_row() for c in sorted(data.categories, key=lambda c: c.sort_order)],
            'imports': [i.to_row() for i in data.imports],
            'allocations': [a.to_row() for a in data.allocations],
        }
        for name, columns in sheet_columns().items():
            ws = wb.create_sheet(name)
            ws.append(list(columns))
            for row_no, row in enumerate(sheet_rows[name], start=2):
                ws.append(row)
                for col_no, value in enumerate(row, start=1):  # '=' 始まりの文字列を数式として保存しない
                    if isinstance(value, str) and value.startswith('='):
                        ws.cell(row=row_no, column=col_no).data_type = 's'

            for idx, col in enumerate(columns, start=1):
                ws.column_dimensions[get_column_letter(idx)].width = EXCEL_COLUMN_WIDTHS.get(col, 14)

            ws.freeze_panes = 'A2'
            if 'amount' in columns:
                for cell in ws[get_column_letter(columns.index('amount') + 1)][1:]:
                    cell.number_format = '#,##0'

        return wb

    def save(self, data: LedgerData) -> Path:
        """一時ファイル → 検証 → 世代バックアップ → 置換 → Dropbox コピーの順で保存する

        Args:
            data: 保存するデータ

        Returns:
            正本のパス
        """
        self.excel_path.parent.mkdir(parents=True, exist_ok=True)
        tmp_path = self.excel_path.with_name(f'~{self.excel_path.stem}.{os.getpid()}.tmp.xlsx')
        try:
            wb = self.build_workbook(data)
            wb.save(tmp_path)
            self.verify(tmp_path)
        except Exception as exc:
            tmp_path.unlink(missing_ok=True)
            raise ExcelSaveError(f'Excel の一時保存に失敗しました: {exc}') from exc

        if self.excel_path.exists():
            self.ensure_writable(tmp_path)
            create_generation_backup(self.excel_path, self.backup_dir, self.backup_generations)

        try:
            os.replace(tmp_path, self.excel_path)
        except PermissionError as exc:
            tmp_path.unlink(missing_ok=True)
            raise ExcelLockedError(self.locked_message()) from exc

        logger.info(
            'excel saved: file=%s transactions=%d rules=%d allocations=%d',
            self.excel_path.name,
            len(data.transactions),
            len(data.merchant_rules),
            len(data.allocations),
        )
        self.last_dropbox_result = self.dropbox.copy(self.excel_path)
        return self.excel_path

    def update(self, mutator: Callable[[LedgerData], Any]) -> LedgerData:
        """読み込み → 変更 → 保存をまとめて行う

        Args:
            mutator: LedgerData を書き換える関数

        Returns:
            保存後の LedgerData
        """
        data = self.load()
        mutator(data)
        self.save(data)
        return data

    # --------------------------------------------------------------- helpers
    def verify(self, path: Path) -> None:
        """保存した一時ファイルを開き直し、必要なシートが揃っているか確認する

        Args:
            path: 一時ファイルのパス
        """
        wb = load_workbook(path, read_only=True)
        try:
            missing = [s for s in sheet_columns() if s not in wb.sheetnames]
            if missing:
                raise ExcelSaveError(f'保存した Excel にシートが不足しています: {missing}')
        finally:
            wb.close()

    def ensure_writable(self, tmp_path: Path) -> None:
        """正本が書き込み可能か確認する（Windows で Excel が開いていると排他ロックされる）

        Args:
            tmp_path: 失敗時に片付ける一時ファイル
        """
        try:
            with open(self.excel_path, 'r+b'):
                pass
        except PermissionError as exc:
            tmp_path.unlink(missing_ok=True)
            raise ExcelLockedError(self.locked_message()) from exc

    def locked_message(self) -> str:
        """ロック時にユーザーへ表示するメッセージ"""
        return (
            f'{self.excel_path.name} が他のアプリ(Excel など)で開かれているため書き込めません。'
            'ファイルを閉じてから再度お試しください。'
        )
