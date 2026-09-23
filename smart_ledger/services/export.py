"""年間表の CSV / Excel エクスポート（集計は aggregation.annual_table の結果をそのまま表にする）"""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook

from ..constants import FORMULA_PREFIXES
from .aggregation import AnnualTable


def annual_table_rows(table: AnnualTable) -> list[list[str | int]]:
    """年間表をヘッダー行付きの 2 次元リストにする（CSV / Excel 共通のレイアウト）

    収入の明細が 1 件も無い年は支出だけの表にして、収入の行を増やさない

    Args:
        table: 年間表

    Returns:
        先頭行が見出し（カテゴリ, 1月〜12月, 年間合計）、続いて支出のカテゴリ行と月間総支出の行。
        収入があればさらに収入のカテゴリ行・月間収入・収支の行が続く
    """
    rows: list[list[str | int]] = [
        ['カテゴリ', *(f'{int(m[5:])}月' for m in table.months), '年間合計'],
        *([r.category, *r.amounts, r.total] for r in table.rows),
        ['月間総支出', *table.monthly_totals, table.total],
    ]
    if not table.income_rows:
        return rows

    rows.extend([r.category, *r.amounts, r.total] for r in table.income_rows)
    rows.append(['月間収入', *table.monthly_incomes, table.income_total])
    rows.append(['収支', *table.monthly_balances, table.balance])
    return rows


def annual_csv(table: AnnualTable, formula_prefixes: str = FORMULA_PREFIXES) -> bytes:
    """年間表を BOM 付き UTF-8 の CSV にする（Windows の Excel で開いても文字化けしない）

    先頭が数式記号のカテゴリ名（旧バージョンで登録できたもの）には ' を付け、Excel で数式として評価させない

    Args:
        table: 年間表
        formula_prefixes: この文字で始まるセルを ' で打ち消す
    """
    rows = [
        [f"'{c}" if isinstance(c, str) and c.startswith(tuple(formula_prefixes)) else c for c in row]
        for row in annual_table_rows(table)
    ]
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return buf.getvalue().encode('utf-8-sig')


def annual_xlsx(table: AnnualTable) -> bytes:
    """年間表を 1 シートの Excel にする（金額は 3 桁区切り、見出し行と見出し列を固定）

    Args:
        table: 年間表
    """
    wb = Workbook()
    wb.remove(wb.worksheets[0])  # 既定シートを捨てて年をタイトルにしたシートだけにする（active は None を返しうる）
    ws = wb.create_sheet(f'{table.year}年')
    for row in annual_table_rows(table):
        ws.append(row)

    for cell in ws['A']:  # 旧バージョンの数式記号始まりを含め、カテゴリ列を数式として保存しない
        cell.data_type = 's'

    ws.freeze_panes = 'B2'
    ws.column_dimensions['A'].width = 18
    for col in ws.iter_cols(min_col=2):
        ws.column_dimensions[col[0].column_letter].width = 12
        for cell in col[1:]:
            cell.number_format = '#,##0'

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
