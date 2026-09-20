"""Excel リポジトリのテスト"""

from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from smart_ledger.constants import DEFAULT_CATEGORIES
from smart_ledger.models import Allocation, ImportRecord, MerchantRule, Transaction
from smart_ledger.services.backup import DropboxBackup
from smart_ledger.services.excel_repository import ExcelRepository, sheet_columns


def test_load_creates_workbook_with_all_sheets_and_categories(repo: ExcelRepository) -> None:
    """正本が無ければ 5 シートと初期カテゴリ入りのファイルを作る

    Args:
        repo: 一時ディレクトリのリポジトリ
    """
    data = repo.load()
    assert repo.excel_path.exists()

    wb = load_workbook(repo.excel_path)
    assert set(sheet_columns()) <= set(wb.sheetnames)
    assert [c.category for c in sorted(data.categories, key=lambda c: c.sort_order)] == list(DEFAULT_CATEGORIES)
    assert [c.value for c in wb['transactions'][1]] == list(Transaction.COLUMNS)
    assert [c.value for c in wb['allocations'][1]] == list(Allocation.COLUMNS)


def test_roundtrip_write_read(repo: ExcelRepository) -> None:
    """全シートの書き込みと読み込みで値が保たれる

    Args:
        repo: 一時ディレクトリのリポジトリ
    """
    data = repo.load()
    data.transactions.append(
        Transaction(
            id='tx_1',
            usage_date=date(2026, 8, 15),
            merchant_raw='ＫＹＡＳＨ',
            merchant_normalized='KYASH',
            amount=10000,
            category='その他',
            confidence=0.55,
            classification_source='jev',
            card='テストカード',
            import_id='imp_1',
            imported_at='2026-09-20 10:00:00',
            row_key='abc',
            memo='メモ',
        )
    )
    data.merchant_rules.append(MerchantRule('KYASH', 'その他', '2026-09-20 10:00:00'))
    data.imports.append(ImportRecord('imp_1', 'a.csv', 'hash', '2026-09-20 10:00:00', 'テストカード', 1))
    data.allocations.append(Allocation('tx_1', '食費', 10000, '内訳'))
    repo.save(data)

    loaded = repo.load()
    tx = loaded.transactions[0]
    assert tx.id == 'tx_1' and tx.usage_date == date(2026, 8, 15) and tx.amount == 10000
    assert tx.merchant_raw == 'ＫＹＡＳＨ' and tx.merchant_normalized == 'KYASH'
    assert tx.confidence == 0.55 and tx.classification_source == 'jev' and tx.memo == 'メモ'
    assert loaded.merchant_rules[0].merchant_pattern == 'KYASH'
    assert loaded.imports[0].file_hash == 'hash'
    assert loaded.allocations[0].amount == 10000


def test_save_creates_generation_backup_and_no_temp_left(repo: ExcelRepository) -> None:
    """2 回目以降の保存で世代バックアップが作られ、一時ファイルは残らない

    Args:
        repo: 一時ディレクトリのリポジトリ
    """
    data = repo.load()
    repo.save(data)
    repo.save(data)
    assert len(list(repo.backup_dir.glob('household_*.xlsx'))) == 2
    assert not list(repo.excel_path.parent.glob('~*.tmp.xlsx'))


def test_backup_pruning(tmp_path: Path) -> None:
    """backup_generations を超えた世代は削除される

    Args:
        tmp_path: pytest の一時ディレクトリ
    """
    repo = ExcelRepository(tmp_path / 'h.xlsx', backup_dir=tmp_path / 'b', backup_generations=2)
    data = repo.load()
    for _ in range(5):
        repo.save(data)

    assert len(list((tmp_path / 'b').glob('h_*.xlsx'))) == 2


def test_dropbox_disabled_is_noop(repo: ExcelRepository) -> None:
    """Dropbox パス未設定ならコピーせず正常終了する

    Args:
        repo: 一時ディレクトリのリポジトリ
    """
    repo.save(repo.load())
    assert repo.last_dropbox_result is None
    assert DropboxBackup(None).enabled is False
    assert DropboxBackup(None).copy(repo.excel_path) is None


def test_dropbox_enabled_copies_latest_and_backup(tmp_path: Path) -> None:
    """Dropbox パス設定時は latest と backup にコピーされる

    Args:
        tmp_path: pytest の一時ディレクトリ
    """
    dropbox = tmp_path / 'Dropbox' / 'SmartLedger'
    repo = ExcelRepository(tmp_path / 'household.xlsx', backup_dir=tmp_path / 'backup', dropbox=DropboxBackup(dropbox))
    repo.save(repo.load())
    assert (dropbox / 'latest' / 'household.xlsx').exists()
    assert len(list((dropbox / 'backup').glob('household_*.xlsx'))) == 1
    assert repo.last_dropbox_result is not None


def test_load_tolerates_missing_columns(tmp_path: Path) -> None:
    """列が少ない旧形式の household.xlsx も読める

    Args:
        tmp_path: pytest の一時ディレクトリ
    """
    path = tmp_path / 'old.xlsx'
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet('transactions')
    ws.append(['id', 'usage_date', 'merchant_raw', 'merchant_normalized', 'amount', 'category'])
    ws.append(['tx_old', '2026-07-01', 'A', 'A', 100, '食費'])
    for name, columns in sheet_columns().items():
        if name != 'transactions':
            wb.create_sheet(name).append(list(columns))

    wb.save(path)
    data = ExcelRepository(path, backup_dir=tmp_path / 'b').load()
    assert data.transactions[0].id == 'tx_old'
    assert data.transactions[0].row_key == '' and data.transactions[0].memo == ''
