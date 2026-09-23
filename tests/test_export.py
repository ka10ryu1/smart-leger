"""年間表の CSV / Excel エクスポートのテスト"""

from __future__ import annotations

import csv
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from smart_ledger.models import Category, LedgerData, Transaction
from smart_ledger.services.aggregation import annual_table
from smart_ledger.services.export import annual_csv, annual_table_rows, annual_xlsx


def sample_ledger(formula_category: str = '=SUM') -> LedgerData:
    """2 か月分の明細と数式記号始まりのカテゴリを含む LedgerData"""
    return LedgerData(
        transactions=[
            Transaction('a', date(2026, 1, 5), 'A', 'A', 1000, '食費'),
            Transaction('b', date(2026, 2, 5), 'B', 'B', 2500, formula_category),
        ],
        categories=[Category('食費', 1), Category(formula_category, 2)],
    )


def test_annual_table_rows_layout() -> None:
    """見出し・カテゴリ行・月間総支出の行が 14 列で並ぶ"""
    rows = annual_table_rows(annual_table(sample_ledger(), 2026))
    assert rows[0] == ['カテゴリ', *(f'{m}月' for m in range(1, 13)), '年間合計']
    assert rows[1] == ['食費', 1000, *([0] * 11), 1000]
    assert rows[-1] == ['月間総支出', 1000, 2500, *([0] * 10), 3500]
    assert all(len(r) == 14 for r in rows)


def test_annual_csv_has_bom_and_values() -> None:
    """CSV は BOM 付き UTF-8 で、Excel 向けに読み戻せる"""
    body = annual_csv(annual_table(sample_ledger(), 2026))
    assert body.startswith(b'\xef\xbb\xbf')  # BOM（不可視文字を直書きしない）
    parsed = list(csv.reader(io.StringIO(body.decode('utf-8-sig'))))
    assert parsed[0][0] == 'カテゴリ' and parsed[0][-1] == '年間合計'
    assert parsed[-1] == ['月間総支出', '1000', '2500', *(['0'] * 10), '3500']


@pytest.mark.parametrize('prefix', ['=', '+', '-', '@'])
def test_annual_csv_neutralizes_formula_like_category(prefix: str) -> None:
    """数式記号始まりの旧カテゴリ名は CSV でも数式にならない

    Args:
        prefix: Excel が数式として解釈しうる先頭文字
    """
    category = f'{prefix}SUM'
    parsed = list(csv.reader(io.StringIO(annual_csv(annual_table(sample_ledger(category), 2026)).decode('utf-8-sig'))))
    assert parsed[2][0] == f"'{category}"  # ' を付けて打ち消す
    assert parsed[1][0] == '食費' and parsed[0][0] == 'カテゴリ'  # 通常の見出し・カテゴリ名は変えない


def test_annual_xlsx_is_readable_and_keeps_formula_like_text() -> None:
    """Excel は openpyxl で開け、'=' 始まりのカテゴリ名が数式にならない"""
    body = annual_xlsx(annual_table(sample_ledger(), 2026))
    ws = load_workbook(io.BytesIO(body)).active
    assert ws.title == '2026年'
    assert ws.freeze_panes == 'B2'
    values = [list(r) for r in ws.iter_rows(values_only=True)]
    assert values[0][0] == 'カテゴリ' and values[-1][0] == '月間総支出'
    assert values[2][0] == '=SUM' and ws.cell(row=3, column=1).data_type == 's'
    assert values[-1][-1] == 3500
    assert ws.cell(row=2, column=2).number_format == '#,##0'


def test_annual_exports_for_empty_year() -> None:
    """明細が無い年も見出しと月間総支出を持つ CSV / Excel を生成する"""
    table = annual_table(LedgerData(), 2027)
    parsed = list(csv.reader(io.StringIO(annual_csv(table).decode('utf-8-sig'))))
    assert parsed == [
        ['カテゴリ', *(f'{m}月' for m in range(1, 13)), '年間合計'],
        ['月間総支出', *(['0'] * 13)],
    ]

    ws = load_workbook(io.BytesIO(annual_xlsx(table))).active
    values = [list(row) for row in ws.iter_rows(values_only=True)]
    assert values == [
        ['カテゴリ', *(f'{m}月' for m in range(1, 13)), '年間合計'],
        ['月間総支出', *([0] * 13)],
    ]


def test_annual_exports_keep_refund_as_negative() -> None:
    """返金は年間合計・CSV・Excelで負数のまま出力する"""
    data = LedgerData(
        transactions=[Transaction('refund', date(2026, 3, 5), '返金', '返金', -500, '食費')],
        categories=[Category('食費', 1)],
    )
    table = annual_table(data, 2026)
    parsed = list(csv.reader(io.StringIO(annual_csv(table).decode('utf-8-sig'))))
    ws = load_workbook(io.BytesIO(annual_xlsx(table))).active

    assert table.monthly_totals[2] == table.total == -500
    assert parsed[1][3] == '-500' and parsed[-1][-1] == '-500'
    assert ws.cell(row=2, column=4).value == -500 and ws.cell(row=3, column=14).value == -500


def test_annual_table_rows_appends_income_block() -> None:
    """収入がある年は、月間総支出の下に収入のカテゴリ行・月間収入・収支が続く"""
    data = sample_ledger()
    data.transactions.append(Transaction('c', date(2026, 2, 5), 'C', 'C', 6000, '売電収入', kind='income'))
    data.categories.append(Category('売電収入', 3))

    rows = annual_table_rows(annual_table(data, 2026))
    assert rows[-4] == ['月間総支出', 1000, 2500, *([0] * 10), 3500]
    assert rows[-3] == ['売電収入', 0, 6000, *([0] * 10), 6000]
    assert rows[-2] == ['月間収入', 0, 6000, *([0] * 10), 6000]
    assert rows[-1] == ['収支', -1000, 3500, *([0] * 10), 2500]
    assert all(len(r) == 14 for r in rows)
